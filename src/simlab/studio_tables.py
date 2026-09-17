"""Bounded spreadsheet parameter import; inspect first, validate before application."""
from __future__ import annotations

import base64
import copy
import csv
import io
import math
import zipfile

from fastapi import HTTPException
from pydantic import BaseModel, ConfigDict, Field

from simlab.manufacturing import ManufacturingConfig

FIELDS = {
    'machines': ['name', 'cycle_time_seconds', 'availability', 'mttr_seconds',
                 'idle_power_kw', 'processing_power_kw'],
    'buffers': ['name', 'capacity', 'delay_seconds'],
}


class TableFile(BaseModel):
    model_config = ConfigDict(extra='forbid')
    filename: str = Field(max_length=200)
    data: str = Field(max_length=1_400_000, repr=False)
    sheet: str | None = Field(default=None, max_length=100)


def read_table(body: TableFile) -> dict:
    try:
        raw = base64.b64decode(body.data, validate=True)
        if len(raw) > 1_000_000:
            raise ValueError('文件不能超过 1 MB。')
        sheets = []
        if body.filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 300 or sum(i.file_size for i in archive.infolist()) > 30_000_000:
                    raise ValueError('工作簿解压后过大，请只保留参数表。')
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheets = book.sheetnames
                if body.sheet and body.sheet not in sheets:
                    raise ValueError('工作表不存在。')
                sheet = book[body.sheet] if body.sheet else book.active
                rows = []
                for index, row in enumerate(sheet.iter_rows()):
                    if index > 500 or len(row) > 40:
                        raise ValueError('最多读取 500 行参数、40 列。')
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError(f'第 {index + 1} 行含公式，请先粘贴为值。')
                    rows.append(['' if cell.value is None else str(cell.value) for cell in row])
            finally:
                book.close()
        elif body.filename.lower().endswith('.csv'):
            try:
                source = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                source = raw.decode('gb18030')
            try:
                dialect = csv.Sniffer().sniff(source[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(io.StringIO(source), dialect))
        else:
            raise ValueError('请选择 .csv 或 .xlsx 文件；旧版 .xls 请先另存为 .xlsx。')
        if len(rows) < 2 or len(rows) > 501:
            raise ValueError('表格须有标题行和 1–500 行数据。')
        headers = [str(value).strip() for value in rows[0]]
        if not headers or len(headers) > 40 or any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('标题行不能为空、重复，且最多 40 列。')
        result = []
        for index, row in enumerate(rows[1:], start=2):
            if not any(str(v).strip() for v in row):
                continue
            if len(row) > len(headers) or any(len(str(value)) > 2000 for value in row):
                raise ValueError(f'第 {index} 行列数或单元格长度超出限制。')
            result.append({'row': index, 'values': dict(zip(headers, row + [''] * (len(headers)-len(row))))})
        if not result:
            raise ValueError('没有可读取的数据行。')
        return {'headers': headers, 'rows': result, 'sheets': sheets}
    except HTTPException:
        raise
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else '无法读取文件，请检查编码或工作簿格式。'
        raise HTTPException(422, message) from None


def map_parameters(config: dict, table: dict, kind: str, mapping: dict, percent: bool) -> dict:
    if kind not in FIELDS or not mapping.get('name') or len(mapping) < 2:
        raise HTTPException(422, '请映射名称及至少一个参数列。')
    if any(key not in FIELDS[kind] or column not in table['headers'] for key,column in mapping.items()):
        raise HTTPException(422, '字段映射无效。')
    if len(set(mapping.values())) != len(mapping):
        raise HTTPException(422, '同一列不能重复映射到多个字段。')
    candidate = copy.deepcopy(config)
    objects = {item['name']: item for item in candidate[kind]}
    errors, changes, seen = [], [], set()
    for row in table['rows']:
        values = row['values']; name = str(values[mapping['name']]).strip()
        if name not in objects and name.startswith("'") and name[1:] in objects:
            name=name[1:]
        if name not in objects or name in seen:
            errors.append({'row':row['row'],'field':'name','message':'名称不存在或在文件中重复；请先在模型中添加对应设备或容器。'})
            continue
        seen.add(name)
        for key,column in mapping.items():
            if key == 'name' or not str(values[column]).strip():
                continue
            try:
                number = float(str(values[column]).strip().rstrip('%'))
                if key == 'availability' and percent: number /= 100
                if not math.isfinite(number): raise ValueError
                if key == 'capacity':
                    if not number.is_integer(): raise ValueError
                    number = int(number)
                old = objects[name][key]
                objects[name][key] = number
                if old != number: changes.append({'row':row['row'],'name':name,'field':key,'before':old,'after':number})
            except (ValueError, TypeError):
                errors.append({'row':row['row'],'field':key,'message':'请输入有限数值；容量须为整数。'})
    if not errors:
        try:
            ManufacturingConfig.model_validate(candidate)
        except ValueError as exc:
            for error in exc.errors():
                loc = error['loc']; name = None
                if len(loc) > 1 and loc[0] == kind and isinstance(loc[1],int): name = candidate[kind][loc[1]]['name']
                row = next((r['row'] for r in table['rows'] if r['values'][mapping['name']].strip() == name), None)
                errors.append({'row':row,'field':'.'.join(map(str,loc)), 'message':error['msg']})
    return {'ok':not errors,'config':candidate if not errors else None,'changes':changes,'errors':errors}
